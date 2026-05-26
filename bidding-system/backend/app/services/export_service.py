"""CSV / 엑셀 내보내기"""
import csv
import io
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.bid_application import BidApplication
from app.models.announcement import Announcement


COLUMNS = [
    ("id", "ID"),
    ("bid_number", "공고번호"),
    ("title", "공고명"),
    ("organization", "발주처"),
    ("category", "업종"),
    ("region", "지역"),
    ("deadline", "마감일"),
    ("bid_price", "투찰가"),
    ("result", "결과"),
    ("result_price", "낙찰가"),
    ("winner_price", "1위 낙찰가"),
    ("our_rank", "순위"),
    ("total_bidders", "총 투찰 업체"),
    ("loss_reason", "유찰 원인"),
    ("submitted_at", "제출일"),
    ("result_updated_at", "결과 업데이트일"),
]


_RESULT_KR = {"won": "낙찰", "lost": "유찰"}


async def _build_rows(db: AsyncSession) -> list[dict]:
    """BidApplication JOIN Announcement — 단일 쿼리"""
    rows_raw = (await db.execute(
        select(
            BidApplication.id,
            BidApplication.bid_price,
            BidApplication.result,
            BidApplication.result_price,
            BidApplication.winner_price,
            BidApplication.our_rank,
            BidApplication.total_bidders,
            BidApplication.loss_reason,
            BidApplication.submitted_at,
            BidApplication.result_updated_at,
            Announcement.bid_number,
            Announcement.title,
            Announcement.organization,
            Announcement.category,
            Announcement.region,
            Announcement.deadline,
        )
        .join(Announcement, BidApplication.announcement_id == Announcement.id, isouter=True)
        .order_by(BidApplication.id)
    )).all()

    return [
        {
            "id": r.id,
            "bid_number": r.bid_number or "",
            "title": r.title or "",
            "organization": r.organization or "",
            "category": r.category or "",
            "region": r.region or "",
            "deadline": r.deadline.strftime("%Y-%m-%d") if r.deadline else "",
            "bid_price": r.bid_price or "",
            "result": _RESULT_KR.get(r.result or "", r.result or ""),
            "result_price": r.result_price or "",
            "winner_price": r.winner_price or "",
            "our_rank": r.our_rank or "",
            "total_bidders": r.total_bidders or "",
            "loss_reason": r.loss_reason or "",
            "submitted_at": r.submitted_at.strftime("%Y-%m-%d %H:%M") if r.submitted_at else "",
            "result_updated_at": r.result_updated_at.strftime("%Y-%m-%d %H:%M") if r.result_updated_at else "",
        }
        for r in rows_raw
    ]


async def export_csv(db: AsyncSession) -> bytes:
    rows = await _build_rows(db)
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=[c[0] for c in COLUMNS])
    writer.writerow({c[0]: c[1] for c in COLUMNS})  # 한글 헤더
    writer.writerows(rows)
    return buf.getvalue().encode("utf-8-sig")  # BOM for Excel 한글


async def export_excel(db: AsyncSession) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    rows = await _build_rows(db)
    wb = Workbook()
    ws = wb.active
    ws.title = "입찰 실적"

    header_fill = PatternFill("solid", fgColor="1E40AF")
    header_font = Font(color="FFFFFF", bold=True)

    headers = [c[1] for c in COLUMNS]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row_idx, row in enumerate(rows, 2):
        for col_idx, (key, _) in enumerate(COLUMNS, 1):
            ws.cell(row=row_idx, column=col_idx, value=row[key])

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
